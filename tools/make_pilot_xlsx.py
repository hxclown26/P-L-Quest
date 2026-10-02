#!/usr/bin/env python3
"""Builds the pilot analysis workbook: the sheet where each person's test scores and survey go, and
the summary that compares before and after against the success criteria of the protocol.

usage: python3 tools/make_pilot_xlsx.py docs/piloto/plantilla-analisis.xlsx
       python3 tools/make_pilot_xlsx.py --key docs/piloto/resultados/clave.json docs/piloto/resultados/plantilla-con-puntaje.xlsx

Sheets: Instrucciones, Datos (one row per participant), Resumen (formulas). The public template holds
no answer key: the scores are entered as 0 (wrong) or 1 (right), as the form export gives them. With
--key (a JSON {"A": [...5 letters], "B": [...5 letters]}) the workbook also gets the sheets Respuestas
(where the answers go as the form exports them, "b) Incentivos" or just "b") and Clave, and scores
them itself; keep that one out of the repository. Needs openpyxl.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from make_review_xlsx import BAD_FILL, BASE_FONT, BORDER, HEADER_FILL, INPUT_FILL, WRAP, style_header

PEOPLE = 30
FIRST, LAST = 2, PEOPLE + 1
GROUPS = ["A→B", "B→A"]
QUESTIONS = 5
GAIN_GOAL = 20          # percentage points of correct answers between before and after
FINISH_GOAL = 0.8       # share of people who finish the game (a half year) ...
FINISH_MINUTES = 30     # ... in at most this many minutes
USEFUL_GOAL = 4         # average usefulness, out of 5

# Datos columns, in order: (header, width, group). The letters below follow from this order.
COLUMNS = (
    [("ID", 7), ("Grupo (forma antes→después)", 14)]
    + [(f"Antes P{i}", 7) for i in range(1, QUESTIONS + 1)] + [("Antes total", 9)]
    + [(f"Después P{i}", 8) for i in range(1, QUESTIONS + 1)] + [("Después total", 9)]
    + [("Ganancia (aciertos)", 11), ("Cambio", 10), ("¿Terminó la partida?", 11), ("Minutos de la partida", 10),
       ("Utilidad (1 a 5)", 10), ("¿Objeción al P&L?", 11), ("Observaciones", 40)]
)
BEFORE = [chr(ord("C") + i) for i in range(QUESTIONS)]          # C..G
BEFORE_TOTAL = "H"
AFTER = [chr(ord("I") + i) for i in range(QUESTIONS)]           # I..M
AFTER_TOTAL, GAIN, CHANGE, FINISHED, MINUTES, USEFUL, OBJECTION = "N", "O", "P", "Q", "R", "S", "T"
COMPLETE = f'">=-{QUESTIONS}"'  # numeric criterion that only matches people with both tests scored


def rng(column):
    return f"Datos!{column}{FIRST}:{column}{LAST}"


ANSWER_BEFORE = [chr(ord("C") + i) for i in range(QUESTIONS)]   # Respuestas!C..G
ANSWER_AFTER = [chr(ord("H") + i) for i in range(QUESTIONS)]    # Respuestas!H..L


def score_formula(row, answer_column, question, post):
    """1 when the first letter of the answer is the key of the form that was taken then, 0 when not."""
    own, other = ("1", "2") if not post else ("2", "1")
    form = f'IF(Datos!$B{row}="A→B",{own},{other})'
    return (f'=IF(OR(Datos!$B{row}="",Respuestas!{answer_column}{row}=""),"",'
            f'IF(LOWER(LEFT(Respuestas!{answer_column}{row},1))=INDEX(Clave!$B$2:$C${QUESTIONS + 1},{question},{form}),1,0))')


def build_answers_and_key(workbook, key):
    answers = workbook.create_sheet("Respuestas")
    headers = ["ID", "Grupo"] + [f"Antes P{i}" for i in range(1, QUESTIONS + 1)] + [f"Después P{i}" for i in range(1, QUESTIONS + 1)]
    for column, header in enumerate(headers, start=1):
        answers.cell(row=1, column=column, value=header)
    for row in range(FIRST, LAST + 1):
        answers.cell(row=row, column=1, value=f"=Datos!A{row}")
        answers.cell(row=row, column=2, value=f"=Datos!B{row}")
        for column in range(3, 3 + 2 * QUESTIONS):
            answers.cell(row=row, column=column).fill = INPUT_FILL
    style_header(answers, [(header, 14) for header in headers])
    clave = workbook.create_sheet("Clave")
    clave["A1"], clave["B1"], clave["C1"] = "Pregunta", "Forma A", "Forma B"
    for question in range(QUESTIONS):
        clave.cell(row=question + 2, column=1, value=question + 1)
        clave.cell(row=question + 2, column=2, value=key["A"][question])
        clave.cell(row=question + 2, column=3, value=key["B"][question])


def build_data(workbook, keyed=False):
    sheet = workbook.create_sheet("Datos")
    for row in range(FIRST, LAST + 1):
        sheet.cell(row=row, column=1, value=f"P{row - 1:02d}")
        if keyed:
            for question in range(QUESTIONS):
                sheet[f"{BEFORE[question]}{row}"] = score_formula(row, ANSWER_BEFORE[question], question + 1, post=False)
                sheet[f"{AFTER[question]}{row}"] = score_formula(row, ANSWER_AFTER[question], question + 1, post=True)
        sheet.cell(row=row, column=8, value=f'=IF(COUNT(C{row}:G{row})={QUESTIONS},SUM(C{row}:G{row}),"")')
        sheet.cell(row=row, column=14, value=f'=IF(COUNT(I{row}:M{row})={QUESTIONS},SUM(I{row}:M{row}),"")')
        sheet.cell(row=row, column=15, value=f'=IF(AND(H{row}<>"",N{row}<>""),N{row}-H{row},"")')
        sheet.cell(row=row, column=16, value=f'=IF(O{row}="","",IF(O{row}>0,"Mejoró",IF(O{row}<0,"Empeoró","Igual")))')
        for column in range(1, len(COLUMNS) + 1):
            cell = sheet.cell(row=row, column=column)
            cell.font, cell.border, cell.alignment = BASE_FONT, BORDER, WRAP
            computed = (1, 8, 14, 15, 16) + ((3, 4, 5, 6, 7, 9, 10, 11, 12, 13) if keyed else ())
            if column not in computed:
                cell.fill = INPUT_FILL
    style_header(sheet, COLUMNS)
    sheet.freeze_panes = "C2"
    lists = [(GROUPS, "B"), (["Sí", "No"], FINISHED), (["Sí", "No"], OBJECTION)]
    for choices, column in lists:
        validation = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"{column}{FIRST}:{column}{LAST}")
    for columns, low, high in (([] if keyed else BEFORE + AFTER, 0, 1), ([USEFUL], 1, 5), ([MINUTES], 0, 600)):
        validation = DataValidation(type="whole", operator="between", formula1=str(low), formula2=str(high), allow_blank=True)
        validation.error = f"Escribe un número entero de {low} a {high}"
        sheet.add_data_validation(validation)
        for column in columns:
            validation.add(f"{column}{FIRST}:{column}{LAST}")
    sheet.conditional_formatting.add(f"{CHANGE}{FIRST}:{CHANGE}{LAST}", FormulaRule(formula=[f'{CHANGE}{FIRST}="Empeoró"'], fill=BAD_FILL))


def mean(column, group=None):
    """Average of a Datos column over the people with both tests scored (optionally one group)."""
    group_part = f',{rng("B")},"{group}"' if group else ""
    return f"IFERROR(AVERAGEIFS({rng(column)},{rng(GAIN)},{COMPLETE}{group_part}),\"\")"


def count(group=None):
    group_part = f',{rng("B")},"{group}"' if group else ""
    return f"COUNTIFS({rng(GAIN)},{COMPLETE}{group_part})"


def pct_points(after, before):
    return f'=IF(OR({after}="",{before}=""),"",({after}-{before})*100)'


def build_summary(workbook):
    sheet = workbook.create_sheet("Resumen")
    sheet.column_dimensions["A"].width = 52
    for column in "BCDEF":
        sheet.column_dimensions[column].width = 15
    bold = Font(name="Calibri", size=10, bold=True)
    title = Font(name="Calibri", size=11, bold=True, color="1F3864")
    sheet["A1"] = "Resumen del piloto"
    sheet["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F3864")

    def put(cell, value, font=BASE_FONT, fmt=None):
        sheet[cell] = value
        sheet[cell].font = font
        if fmt:
            sheet[cell].number_format = fmt

    # --- learning
    put("A3", "Aprendizaje (test antes y después)", title)
    put("A4", "Personas con las dos pruebas completas")
    put("B4", f"={count()}", bold)
    put("A5", "Promedio antes (% de aciertos)")
    put("B5", f'=IF(B4=0,"",{mean(BEFORE_TOTAL)}/{QUESTIONS})', bold, "0.0%")
    put("A6", "Promedio después (% de aciertos)")
    put("B6", f'=IF(B4=0,"",{mean(AFTER_TOTAL)}/{QUESTIONS})', bold, "0.0%")
    put("A7", "Ganancia (puntos porcentuales)")
    put("B7", pct_points("B6", "B5"), bold, "0.0")
    put("A8", f"Meta: +{GAIN_GOAL} puntos porcentuales")
    put("B8", f'=IF(B7="","",IF(B7>={GAIN_GOAL},"Cumple","No cumple"))', bold)
    put("A9", "Mejoraron / Igual / Empeoraron")
    for column, label in zip("BCD", ["Mejoró", "Igual", "Empeoró"]):
        put(f"{column}9", f'=COUNTIF({rng(CHANGE)},"{label}")', bold)
    put("A10", "Prueba de signos, p bilateral (descriptiva: n es pequeño)")
    put("B10", '=IF(B9+D9=0,"",MIN(1,2*(1-BINOMDIST(MAX(B9,D9)-1,B9+D9,0.5,TRUE))))', bold, "0.000")

    # --- per question
    put("A12", "Por pregunta (la que menos subió pide revisar el juego o la pregunta)", title)
    for column, header in zip("BCD", ["Antes", "Después", "Ganancia (pp)"]):
        put(f"{column}13", header, bold)
    for offset in range(QUESTIONS):
        row = 14 + offset
        put(f"A{row}", f"Pregunta {offset + 1}")
        put(f"B{row}", f'=IF($B$4=0,"",{mean(BEFORE[offset])})', BASE_FONT, "0.0%")
        put(f"C{row}", f'=IF($B$4=0,"",{mean(AFTER[offset])})', BASE_FONT, "0.0%")
        put(f"D{row}", pct_points(f"C{row}", f"B{row}"), BASE_FONT, "0.0")

    # --- per group (a big gap between groups points at forms of different difficulty)
    put("A20", "Por grupo (si los dos difieren mucho, las formas A y B no pesan igual)", title)
    for column, header in zip("BCDE", ["Personas", "Antes", "Después", "Ganancia (pp)"]):
        put(f"{column}21", header, bold)
    for offset, group in enumerate(GROUPS):
        row = 22 + offset
        put(f"A{row}", f"Grupo {group}")
        put(f"B{row}", f"={count(group)}")
        put(f"C{row}", f'=IF(B{row}=0,"",{mean(BEFORE_TOTAL, group)}/{QUESTIONS})', BASE_FONT, "0.0%")
        put(f"D{row}", f'=IF(B{row}=0,"",{mean(AFTER_TOTAL, group)}/{QUESTIONS})', BASE_FONT, "0.0%")
        put(f"E{row}", pct_points(f"D{row}", f"C{row}"), BASE_FONT, "0.0")

    # --- the game itself
    put("A25", "El juego (observación y encuesta)", title)
    put("A26", "Personas observadas (marcaron si terminaron la partida)")
    put("B26", f'=COUNTIF({rng(FINISHED)},"Sí")+COUNTIF({rng(FINISHED)},"No")', bold)
    put("A27", f"Terminaron la partida en {FINISH_MINUTES} minutos o menos")
    put("B27", f'=IF(B26=0,"",COUNTIFS({rng(FINISHED)},"Sí",{rng(MINUTES)},"<={FINISH_MINUTES}")/B26)', bold, "0.0%")
    put("A28", f"Meta: al menos {int(FINISH_GOAL * 100)}%")
    put("B28", f'=IF(B27="","",IF(B27>={FINISH_GOAL},"Cumple","No cumple"))', bold)
    put("A29", "Utilidad promedio (1 a 5)")
    put("B29", f'=IFERROR(AVERAGE({rng(USEFUL)}),"")', bold, "0.00")
    put("A30", f"Meta: al menos {USEFUL_GOAL}")
    put("B30", f'=IF(B29="","",IF(B29>={USEFUL_GOAL},"Cumple","No cumple"))', bold)
    put("A31", "Personas de finanzas con objeción al P&L en pantalla")
    put("B31", f'=COUNTIF({rng(OBJECTION)},"Sí")', bold)
    put("A32", "Meta: ninguna")
    put("B32", '=IF(B26=0,"",IF(B31=0,"Cumple","No cumple"))', bold)

    # --- decision
    put("A34", "Decisión", title)
    put("B34", '=IF(OR(B8="",B28="",B30="",B32=""),"Faltan datos",IF(AND(B8="Cumple",B28="Cumple",B30="Cumple",B32="Cumple"),"Evaluar Demo 5","Iterar con lo que mostró el piloto"))', bold)
    sheet.page_setup.orientation = "portrait"
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0


def build_instructions(workbook, keyed=False):
    sheet = workbook.active
    sheet.title = "Instrucciones"
    sheet.column_dimensions["A"].width = 100
    lines = [
        ("Plantilla de análisis del piloto de P&L Quest", Font(name="Calibri", size=14, bold=True, color="1F3864")),
        ("1. En la hoja Datos hay una fila por persona (hasta 30). Las celdas amarillas se llenan; las demás se calculan.", BASE_FONT),
        ("2. Test antes y después: escribe 1 si acertó la pregunta y 0 si no (el formulario lo entrega como puntaje por pregunta). Elige el grupo: A→B hizo la forma A antes y la B después; B→A al revés.", BASE_FONT),
        ("3. Observación: ¿terminó la partida? y en cuántos minutos (el juego los muestra al final, junto al plan: «mm:ss min»). Encuesta: utilidad de 1 a 5 y si una persona de finanzas objetó el P&L en pantalla.", BASE_FONT),
        ("4. La hoja Resumen compara antes y después con las metas del protocolo y propone una decisión. Con 12 a 15 personas el resultado es descriptivo, no una prueba estadística.", BASE_FONT),
        ("5. Esta plantilla no trae la clave de respuestas. Los datos de las personas no se guardan en el repositorio: usa IDs (P01, P02...) y guarda el archivo completado fuera de él.", BASE_FONT),
    ]
    if keyed:
        lines[2] = ("2. Test antes y después: en la hoja Respuestas pega lo que respondió cada persona tal como lo exporta el formulario («b) Incentivos» o solo «b»); la hoja Datos lo puntúa sola con la hoja Clave. Elige el grupo en Datos: A→B hizo la forma A antes y la B después; B→A al revés.", BASE_FONT)
        lines[5] = ("5. Esta versión TRAE la clave de respuestas (hoja Clave): no la subas a GitHub ni la compartas con quienes responden. Guárdala en docs/piloto/resultados/, que git ignora.", BASE_FONT)
    for row, (text, font) in enumerate(lines, start=1):
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font, cell.alignment = font, WRAP
        sheet.row_dimensions[row].height = 30 if row > 1 else 22


def main(path, key_path=None):
    key = json.load(open(key_path)) if key_path else None
    workbook = Workbook()
    build_instructions(workbook, keyed=key is not None)
    build_data(workbook, keyed=key is not None)
    build_summary(workbook)
    if key is not None:
        build_answers_and_key(workbook, key)
    workbook.save(path)
    print(f"{path}: {PEOPLE} rows for participants" + (", scored from the key" if key else ""))


if __name__ == "__main__":
    args = sys.argv[1:]
    if len(args) == 3 and args[0] == "--key":
        main(args[2], args[1])
    elif len(args) == 1:
        main(args[0])
    else:
        sys.exit(__doc__)
