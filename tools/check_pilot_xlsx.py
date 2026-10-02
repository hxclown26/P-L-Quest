#!/usr/bin/env python3
"""Rehearses the pilot analysis workbook with synthetic participants and checks its Resumen sheet.

usage: python3 tools/check_pilot_xlsx.py docs/piloto/plantilla-analisis.xlsx

Fills copies with 14 made-up participants (fixed seeds, one of them with only the "before" test;
one rehearsal where every goal is met), has LibreOffice recalculate them and compares every number
of the Resumen with the same figure taken in Python. The blank template must show "Faltan datos"
and no errors. The variant that scores the raw answers from a key is rehearsed with a made-up key
and answers written the way a form exports them. Needs openpyxl and soffice.
"""
import json
import math
import random
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

QUESTIONS, PEOPLE = 5, 14
GOALS = {"gain": 20, "finish": 0.8, "minutes": 30, "useful": 4}


def synthetic(seed, goals_met=False):
    rng = random.Random(seed)
    people = []
    for index in range(PEOPLE):
        before_p, after_p = [rng.uniform(0.3, 0.55) for _ in range(QUESTIONS)], [rng.uniform(0.6, 0.9) for _ in range(QUESTIONS)]
        people.append({
            "group": "A→B" if index % 2 == 0 else "B→A",
            "before": [int(rng.random() < p) for p in before_p],
            "after": [int(rng.random() < p) for p in after_p] if index != PEOPLE - 1 else None,
            "finished": "Sí" if goals_met or rng.random() < 0.9 else "No",
            "minutes": rng.randint(18, 28 if goals_met else 40),
            "useful": rng.randint(4 if goals_met else 3, 5),
            "objection": "No" if goals_met else ("Sí" if rng.random() < 0.1 else "No"),
        })
    return people


def write(path, people):
    book = load_workbook(path)
    sheet = book["Datos"]
    for offset, person in enumerate(people):
        row = offset + 2
        sheet.cell(row=row, column=2, value=person["group"])
        for i, value in enumerate(person["before"]):
            sheet.cell(row=row, column=3 + i, value=value)
        for i, value in enumerate(person["after"] or []):
            sheet.cell(row=row, column=9 + i, value=value)
        for column, key in ((17, "finished"), (18, "minutes"), (19, "useful"), (20, "objection")):
            sheet.cell(row=row, column=column, value=person[key])
    book.save(path)


LETTERS = "abcdefg"
FAKE_KEY = {"A": list("bdgcc"), "B": list("begbd")}


def write_keyed(path, people):
    """Writes the group and the observations into Datos, and the raw answers into Respuestas."""
    book = load_workbook(path)
    data, answers = book["Datos"], book["Respuestas"]
    rng = random.Random(21)
    for offset, person in enumerate(people):
        row = offset + 2
        data.cell(row=row, column=2, value=person["group"])
        for column, key in ((17, "finished"), (18, "minutes"), (19, "useful"), (20, "objection")):
            data.cell(row=row, column=column, value=person[key])
        forms = ("A", "B") if person["group"] == "A→B" else ("B", "A")
        for phase, form, scores, first in (("before", forms[0], person["before"], 3), ("after", forms[1], person["after"], 8)):
            if scores is None:
                continue
            for i, ok in enumerate(scores):
                right = FAKE_KEY[form][i]
                letter = right if ok else rng.choice([c for c in LETTERS if c != right])
                text = f"{letter}) una opción" if i % 2 == 0 else letter.upper()   # as exported, or just the letter
                answers.cell(row=row, column=first + i, value=text)
    book.save(path)


def mean(values):
    return sum(values) / len(values) if values else None


def expected(people):
    done = [p for p in people if p["after"] is not None]
    before_total, after_total = [sum(p["before"]) for p in done], [sum(p["after"]) for p in done]
    gains = [a - b for a, b in zip(after_total, before_total)]
    better, same, worse = sum(g > 0 for g in gains), sum(g == 0 for g in gains), sum(g < 0 for g in gains)
    n, k = better + worse, max(better, worse)
    p_value = min(1.0, 2 * sum(math.comb(n, i) for i in range(k, n + 1)) / 2 ** n) if n else None
    gain = (mean(after_total) - mean(before_total)) / QUESTIONS * 100
    cells = {"B4": len(done), "B5": mean(before_total) / QUESTIONS, "B6": mean(after_total) / QUESTIONS, "B7": gain,
             "B8": "Cumple" if gain >= GOALS["gain"] else "No cumple", "B9": better, "C9": same, "D9": worse, "B10": p_value}
    for i in range(QUESTIONS):
        before, after = mean([p["before"][i] for p in done]), mean([p["after"][i] for p in done])
        cells.update({f"B{14 + i}": before, f"C{14 + i}": after, f"D{14 + i}": (after - before) * 100})
    for row, group in ((22, "A→B"), (23, "B→A")):
        members = [p for p in done if p["group"] == group]
        before, after = mean([sum(p["before"]) for p in members]) / QUESTIONS, mean([sum(p["after"]) for p in members]) / QUESTIONS
        cells.update({f"B{row}": len(members), f"C{row}": before, f"D{row}": after, f"E{row}": (after - before) * 100})
    finish = sum(p["finished"] == "Sí" and p["minutes"] <= GOALS["minutes"] for p in people) / len(people)
    useful = mean([p["useful"] for p in people])
    objections = sum(p["objection"] == "Sí" for p in people)
    cells.update({"B26": len(people), "B27": finish, "B28": "Cumple" if finish >= GOALS["finish"] else "No cumple",
                  "B29": useful, "B30": "Cumple" if useful >= GOALS["useful"] else "No cumple", "B31": objections,
                  "B32": "Cumple" if objections == 0 else "No cumple"})
    ok = all(cells[c] == "Cumple" for c in ("B8", "B28", "B30", "B32"))
    cells["B34"] = "Evaluar Demo 5" if ok else "Iterar con lo que mostró el piloto"
    return cells


def recalculate(path, folder):
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(Path(folder) / "out"), str(path)],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(Path(folder) / "out" / Path(path).name, data_only=True)["Resumen"]


def differences(summary, want):
    wrong = []
    for cell, value in want.items():
        got = summary[cell].value
        same = (abs(got - value) <= 1e-9) if isinstance(value, float) and isinstance(got, (int, float)) else got == value
        if not same:
            wrong.append((cell, value, got))
    return wrong


def main(path):
    book = load_workbook(path)
    assert book.sheetnames == ["Instrucciones", "Datos", "Resumen"], book.sheetnames
    assert book["Datos"].max_row == 31
    with tempfile.TemporaryDirectory() as folder:
        blank = Path(folder) / "blank.xlsx"
        shutil.copy(path, blank)
        summary = recalculate(blank, folder)
        errors = [c.coordinate for row in summary.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("#")]
        assert not errors and summary["B34"].value == "Faltan datos", (errors, summary["B34"].value)
        print("ok blank template: no errors, decision: Faltan datos")
        for label, seed, goals_met in (("seed 3", 3, False), ("seed 11", 11, False), ("goals met", 5, True)):
            copy = Path(folder) / f"rehearsal{seed}.xlsx"
            shutil.copy(path, copy)
            people = synthetic(seed, goals_met)
            write(copy, people)
            want = expected(people)
            wrong = differences(recalculate(copy, folder), want)
            if wrong:
                sys.exit(f"{label}: Resumen differs from the Python figures: {wrong}")
            print(f"ok {label}: {len(want)} cells match; gain {want['B7']:.1f} pp; decision: {want['B34']}")
            if goals_met:
                assert want["B34"] == "Evaluar Demo 5", want["B34"]
        # the variant that scores the answers itself, from a made-up key
        key_file = Path(folder) / "key.json"
        key_file.write_text(json.dumps(FAKE_KEY))
        keyed = Path(folder) / "keyed.xlsx"
        script = Path(__file__).with_name("make_pilot_xlsx.py")
        subprocess.run([sys.executable, str(script), "--key", str(key_file), str(keyed)], check=True, capture_output=True)
        people = synthetic(8)
        write_keyed(keyed, people)
        want = expected(people)
        wrong = differences(recalculate(keyed, folder), want)
        if wrong:
            sys.exit(f"keyed: Resumen differs from the Python figures: {wrong}")
        print(f"ok keyed: {len(want)} cells match from raw answers scored against the key; gain {want['B7']:.1f} pp")


if __name__ == "__main__":
    main(sys.argv[1])
