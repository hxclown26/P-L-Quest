#!/usr/bin/env python3
"""Builds the domain-review workbook from the rows of tools/export-review.js.

usage: node tools/export-review.js | python3 tools/make_review_xlsx.py docs/revision/revision-dominio.xlsx

Sheets: Instrucciones, Revisión (one row per answer), Problemas (one row per problem), Resumen
(formulas over the two reviewer sheets). Needs openpyxl.
"""
import json
import math
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

REVIEW = "Revisión"
PROBLEMS = "Problemas"
LINES = ["Ventas", "Incentivos", "Costo", "Flete", "Direct Chg", "SG&A"]
LINE_CHOICES = LINES + ["Fuera del P&L"]
VERDICTS = ["Sí", "No", "Ambiguo"]
VOICES = ["CLIENTE", "PLANTA", "ENTORNO", "ESTRATEGIA"]
TYPES = ["Equilibrada", "Atajo", "Ceder", "Pasiva"]
APPROVAL_SHARE = 0.9

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")
WARN_FILL = PatternFill("solid", fgColor="FFE699")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE_FONT = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

# (header, width)
REVIEW_COLUMNS = [
    ("ID", 7), ("Meses", 11), ("Voz", 12), ("Segmento", 12), ("Escena", 38), ("Respuesta", 20), ("Descripción", 52),
    ("Efecto en el P&L", 28), ("Línea que mueve", 14), ("Intención del autor", 14), ("Después de responder", 52),
    ("¿Realista?", 12), ("Línea correcta", 14), ("Comentario", 40),
    ("Scene (EN)", 38), ("Answer (EN)", 20), ("Description (EN)", 52), ("Effect on the P&L (EN)", 28), ("After answering (EN)", 52),
    ("¿Traducción correcta?", 14),
]
PROBLEM_COLUMNS = [
    ("ID", 7), ("Meses", 11), ("Voz", 12), ("Segmento", 12), ("Título", 24), ("Escena", 44), ("Ficha completa", 60),
    ("Respuesta equilibrada", 60),
    ("¿Escena realista?", 14), ("¿La equilibrada es la mejor?", 16), ("¿Falta una opción mejor? ¿Cuál?", 40), ("Comentario", 40),
]
# The letters of the columns the formulas read and of the ones the reviewer fills in.
REV_VOICE, REV_LINE, REV_TYPE = "C", "I", "J"
REV_REALISTIC, REV_CORRECT_LINE, REV_TRANSLATION = "L", "M", "T"
REV_INPUTS = {12, 13, 14, 20}
PROB_SCENE, PROB_BEST = "I", "J"
PROB_INPUTS = {9, 10, 11, 12}


def months_label(months):
    return "Cualquiera" if months == "any" else f"{months.split('-')[0]} a {months.split('-')[1]}"


def style_header(sheet, columns):
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_title_rows = "1:1"


def fit_row(sheet, row, columns):
    lines = 1
    for index, (_, width) in enumerate(columns, start=1):
        value = sheet.cell(row=row, column=index).value
        if isinstance(value, str):
            lines = max(lines, math.ceil(len(value) / max(width - 2, 1)))
    sheet.row_dimensions[row].height = 13.5 * lines


def write_row(sheet, row, values, columns, input_columns):
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.font = BASE_FONT
        cell.alignment = WRAP
        cell.border = BORDER
        if index in input_columns:
            cell.fill = INPUT_FILL
    fit_row(sheet, row, columns)


def add_list(sheet, choices, column, first, last):
    validation = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
    validation.error = "Elige un valor de la lista"
    validation.errorTitle = "Valor no válido"
    sheet.add_data_validation(validation)
    validation.add(f"{column}{first}:{column}{last}")


def build_review(workbook, rows):
    sheet = workbook.create_sheet(REVIEW)
    for offset, row in enumerate(rows, start=2):
        es, en = row["es"], row["en"]
        values = [
            row["id"], months_label(row["months"]), es["voice"], es["segment"], es["scene"], es["name"], es["desc"],
            es["effect"], es["line"], es["type"], es["story"], None, None, None,
            en["scene"], en["name"], en["desc"], en["effect"], en["story"], None,
        ]
        write_row(sheet, offset, values, REVIEW_COLUMNS, REV_INPUTS)
    last = len(rows) + 1
    style_header(sheet, REVIEW_COLUMNS)
    r, c, t, line = REV_REALISTIC, REV_CORRECT_LINE, REV_TRANSLATION, REV_LINE
    add_list(sheet, VERDICTS, r, 2, last)
    add_list(sheet, LINE_CHOICES, c, 2, last)
    add_list(sheet, ["Sí", "No"], t, 2, last)
    sheet.conditional_formatting.add(f"{r}2:{r}{last}", FormulaRule(formula=[f'{r}2="No"'], fill=BAD_FILL))
    sheet.conditional_formatting.add(f"{r}2:{r}{last}", FormulaRule(formula=[f'{r}2="Ambiguo"'], fill=WARN_FILL))
    sheet.conditional_formatting.add(f"{c}2:{c}{last}", FormulaRule(formula=[f'AND({c}2<>"",{c}2<>{line}2)'], fill=BAD_FILL))
    sheet.conditional_formatting.add(f"{t}2:{t}{last}", FormulaRule(formula=[f'{t}2="No"'], fill=BAD_FILL))
    return last


def build_problems(workbook, rows):
    sheet = workbook.create_sheet(PROBLEMS)
    balanced = [row for row in rows if row["type"] == "smart"]
    for offset, row in enumerate(balanced, start=2):
        es = row["es"]
        values = [
            row["id"], months_label(row["months"]), es["voice"], es["segment"], es["title"], es["scene"],
            f'{es["brief"]} ({es["facts"]})', f'{es["name"]}: {es["desc"]}', None, None, None, None,
        ]
        write_row(sheet, offset, values, PROBLEM_COLUMNS, PROB_INPUTS)
    last = len(balanced) + 1
    style_header(sheet, PROBLEM_COLUMNS)
    add_list(sheet, VERDICTS, PROB_SCENE, 2, last)
    add_list(sheet, VERDICTS, PROB_BEST, 2, last)
    sheet.conditional_formatting.add(f"{PROB_SCENE}2:{PROB_BEST}{last}", FormulaRule(formula=[f'{PROB_SCENE}2="No"'], fill=BAD_FILL))
    return last


def build_instructions(workbook):
    sheet = workbook.active
    sheet.title = "Instrucciones"
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 100
    lines = [
        ("P&L Quest: ficha de revisión de dominio", None),
        ("Para qué sirve", "Comprobar que las 48 situaciones del modo Año completo (192 respuestas) suenan reales en el negocio y que cada respuesta mueve las líneas correctas del P&L, en el sentido correcto. Los clientes son hoteles, hospitales, industria de alimentos e industria en general. Los datos del juego son ficticios."),
        ("Cuánto toma", "Entre 1 y 1,5 horas. Lee una fila a la vez: la escena, la respuesta, lo que el juego dice que hace en el P&L y lo que cuenta después de responder."),
        ("Hoja Revisión", "Una fila por respuesta. «Efecto en el P&L» muestra las líneas que se mueven y hacia dónde (↑ la línea sube, ↓ baja; una respuesta de crecimiento trae además su volumen y su precio, en % de las ventas del plan). «Después de responder» es la historia que el juego cuenta. Completa las celdas amarillas: ¿Realista? (Sí, No o Ambiguo), Línea correcta (la del P&L que a tu juicio mueve más la respuesta; «Fuera del P&L» si es caja, deuda, impuestos u otra cosa) y Comentario (obligatorio si marcas No, Ambiguo o una línea distinta). Las columnas en inglés son para quien revise la traducción."),
        ("Hoja Problemas", "Una fila por situación, con su ficha completa: lo que la persona lee antes de decidir. Dime si la escena es realista, si la respuesta equilibrada es de verdad la mejor y si falta una opción mejor."),
        ("Líneas rojas y crecimiento", "Seis respuestas son líneas rojas (incumplir un contrato, inflar una cifra, ocultar un incidente): el juego las castiga con una multa y no hay resultado que las compre de vuelta; marca si de verdad lo son. Cuatro situaciones de crecimiento (m06c, m08c, m09c, m12c) y dos mixtas (m03s, m06s) se miden en volumen y precio: revisa si el trueque entre vender más y ganar menos es creíble."),
        ("Hoja Resumen", "Se calcula sola con lo que completes. No hay que tocarla."),
        ("Las 6 líneas", "Ventas: ventas antes de incentivos. Incentivos: rebates y descuentos a clientes. Costo: costo de ventas (insumos, producción, mantención). Flete: llevar el producto al cliente. Direct Chg: cargos directos al cliente, como el comodato de equipos o las visitas de servicio. SG&A: ventas, generales y administración. Debajo de la línea de SG&A el juego termina en el OI: caja, capital de trabajo, deuda, intereses e impuestos quedan fuera."),
        ("Los 4 tipos de respuesta", "Equilibrada: la que el juego premia (mejora el OI con una contrapartida bien diseñada). Atajo: sube el OI hoy y cobra después. Ceder: cede a la voz en tensión y el OI paga. Pasiva: no hacer nada."),
        ("Criterio de aprobación", f"Al menos {int(APPROVAL_SHARE * 100)}% de las 192 respuestas marcadas «Sí» en realista y ninguna línea distinta sin resolver. Lo que marques No, Ambiguo o con otra línea se corrige y se vuelve a revisar."),
        ("Revisor", None),
        ("Nombre", None),
        ("Área (comercial, operaciones, otra)", None),
        ("Fecha", None),
    ]
    for row, (label, text) in enumerate(lines, start=1):
        head = sheet.cell(row=row, column=1, value=label)
        body = sheet.cell(row=row, column=2, value=text)
        head.font = Font(name="Calibri", size=14 if row == 1 else 10, bold=True, color="1F3864")
        body.font = BASE_FONT
        head.alignment = body.alignment = WRAP
        if row == 1:
            head.alignment = Alignment(vertical="top")
        if text:
            sheet.row_dimensions[row].height = 13.5 * max(1, math.ceil(len(text) / 95))
        if row > len(lines) - 3:
            body.fill = INPUT_FILL
            body.border = BORDER
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


SUMMARY_ROW = {"total": 3, "reviewed": 4, "yes": 5, "no": 6, "ambiguous": 7, "share": 8, "lines": 9,
               "translations": 10, "scenes": 11, "best": 12, "result": 13}


def build_summary(workbook, review_last, problems_last):
    sheet = workbook.create_sheet("Resumen")
    sheet.column_dimensions["A"].width = 44
    for column in "BCDE":
        sheet.column_dimensions[column].width = 14
    rev, prob, row = f"'{REVIEW}'", f"'{PROBLEMS}'", SUMMARY_ROW
    realistic = lambda verdict: f"COUNTIF({rev}!{REV_REALISTIC}2:{REV_REALISTIC}{review_last},\"{verdict}\")"
    total = f"COUNTA({rev}!A2:A{review_last})"
    cells = [
        ("total", "Respuestas en la ficha", f"={total}"),
        ("reviewed", "Respuestas revisadas (realista marcado)", f"=COUNTA({rev}!{REV_REALISTIC}2:{REV_REALISTIC}{review_last})"),
        ("yes", "Realista: Sí", f"={realistic('Sí')}"),
        ("no", "Realista: No", f"={realistic('No')}"),
        ("ambiguous", "Realista: Ambiguo", f"={realistic('Ambiguo')}"),
        ("share", "% realista (Sí / total)", f"=IF({total}=0,0,{realistic('Sí')}/{total})"),
        ("lines", "Líneas distintas a la del juego", f"=SUMPRODUCT(({rev}!{REV_CORRECT_LINE}2:{REV_CORRECT_LINE}{review_last}<>\"\")*({rev}!{REV_CORRECT_LINE}2:{REV_CORRECT_LINE}{review_last}<>{rev}!{REV_LINE}2:{REV_LINE}{review_last}))"),
        ("translations", "Traducciones marcadas «No»", f"=COUNTIF({rev}!{REV_TRANSLATION}2:{REV_TRANSLATION}{review_last},\"No\")"),
        ("scenes", "Escenas no realistas", f"=COUNTIF({prob}!{PROB_SCENE}2:{PROB_SCENE}{problems_last},\"No\")"),
        ("best", "Equilibradas que no son la mejor", f"=COUNTIF({prob}!{PROB_BEST}2:{PROB_BEST}{problems_last},\"No\")"),
        ("result", "Resultado", f"=IF(AND(B{row['share']}>={APPROVAL_SHARE},B{row['lines']}=0),\"Aprobado\",\"Revisar\")"),
    ]
    sheet["A1"] = "Resumen de la revisión"
    sheet["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F3864")
    for key, label, formula in cells:
        sheet.cell(row=row[key], column=1, value=label).font = BASE_FONT
        cell = sheet.cell(row=row[key], column=2, value=formula)
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.border = BORDER
    sheet.cell(row=row["share"], column=2).number_format = "0.0%"
    break_down(sheet, 16, "Realista por voz", VOICES, REV_VOICE, review_last)
    break_down(sheet, 23, "Realista por tipo de respuesta", TYPES, REV_TYPE, review_last)


def break_down(sheet, start, title, names, column, last):
    rev = f"'{REVIEW}'"
    sheet.cell(row=start, column=1, value=title).font = Font(name="Calibri", size=10, bold=True, color="1F3864")
    for column_index, header in enumerate(["Respuestas", "Sí", "No", "Ambiguo"], start=2):
        sheet.cell(row=start, column=column_index, value=header).font = Font(name="Calibri", size=10, bold=True)
    for offset, name in enumerate(names, start=start + 1):
        sheet.cell(row=offset, column=1, value=name).font = BASE_FONT
        sheet.cell(row=offset, column=2, value=f"=COUNTIF({rev}!{column}2:{column}{last},\"{name}\")")
        for column_index, verdict in enumerate(VERDICTS, start=3):
            sheet.cell(row=offset, column=column_index, value=f"=COUNTIFS({rev}!{column}2:{column}{last},\"{name}\",{rev}!{REV_REALISTIC}2:{REV_REALISTIC}{last},\"{verdict}\")")


def main(path):
    rows = json.load(sys.stdin)
    workbook = Workbook()
    build_instructions(workbook)
    review_last = build_review(workbook, rows)
    problems_last = build_problems(workbook, rows)
    build_summary(workbook, review_last, problems_last)
    workbook.save(path)
    print(f"{path}: {len(rows)} answers, {problems_last - 1} problems")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
