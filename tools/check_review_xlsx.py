#!/usr/bin/env python3
"""Rehearses the review workbook with synthetic answers and checks its Resumen sheet.

usage: python3 tools/check_review_xlsx.py docs/revision/revision-dominio.xlsx

Fills a copy with made-up verdicts (fixed seed), has LibreOffice recalculate it and compares every
number of the Resumen sheet with the same count taken in Python. Needs openpyxl and soffice.
"""
import random
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

REVIEW, PROBLEMS = "Revisión", "Problemas"
VERDICTS = ["Sí", "No", "Ambiguo"]
LINES = ["Ventas", "Incentivos", "Costo", "Flete", "Direct Chg", "SG&A", "Fuera del P&L"]


def fill(path, seed=7):
    """Writes made-up answers into the reviewer columns and returns the counts they imply."""
    rng = random.Random(seed)
    book = load_workbook(path)
    review, problems = book[REVIEW], book[PROBLEMS]
    tally = Counter()
    for row in range(2, review.max_row + 1):
        verdict = rng.choices(VERDICTS, weights=[80, 8, 12])[0]
        line = review.cell(row=row, column=7).value
        other = rng.choice([l for l in LINES if l != line]) if rng.random() < 0.06 else None
        translation = "No" if rng.random() < 0.04 else "Sí"
        review.cell(row=row, column=9, value=verdict)
        review.cell(row=row, column=10, value=other)
        review.cell(row=row, column=15, value=translation)
        tally["total"] += 1
        tally[f"yes_{verdict}"] += 1
        tally["lines"] += other is not None
        tally["translations"] += translation == "No"
        voice, kind = review.cell(row=row, column=3).value, review.cell(row=row, column=8).value
        tally[("voice", voice, verdict)] += 1
        tally[("type", kind, verdict)] += 1
    for row in range(2, problems.max_row + 1):
        scene, best = rng.choices(VERDICTS, weights=[85, 5, 10])[0], rng.choices(VERDICTS, weights=[85, 5, 10])[0]
        problems.cell(row=row, column=7, value=scene)
        problems.cell(row=row, column=8, value=best)
        tally["scenes"] += scene == "No"
        tally["best"] += best == "No"
    book.save(path)
    return tally


def recalculate(path, out_dir):
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out_dir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(Path(out_dir) / Path(path).name, data_only=True)


def expected_cells(tally):
    share = tally["yes_Sí"] / tally["total"]
    cells = {"B3": tally["total"], "B4": tally["total"], "B5": tally["yes_Sí"], "B6": tally["yes_No"],
             "B7": tally["yes_Ambiguo"], "B8": share, "B9": tally["lines"], "B10": tally["translations"],
             "B11": tally["scenes"], "B12": tally["best"],
             "B13": "Aprobado" if share >= 0.9 and tally["lines"] == 0 else "Revisar"}
    for start, kind, names in ((16, "voice", ["CLIENTE", "PLANTA", "ENTORNO", "ESTRATEGIA"]),
                               (23, "type", ["Equilibrada", "Atajo", "Ceder", "Pasiva"])):
        for offset, name in enumerate(names, start=start + 1):
            counts = [tally[(kind, name, verdict)] for verdict in VERDICTS]
            cells[f"B{offset}"] = sum(counts)
            for column, count in zip("CDE", counts):
                cells[f"{column}{offset}"] = count
    return cells


def main(path):
    book = load_workbook(path)
    assert book.sheetnames == ["Instrucciones", REVIEW, PROBLEMS, "Resumen"], book.sheetnames
    assert book[REVIEW].max_row == 193 and book[PROBLEMS].max_row == 49
    assert len(book[REVIEW].data_validations.dataValidation) == 3 and len(book[PROBLEMS].data_validations.dataValidation) == 2
    with tempfile.TemporaryDirectory() as folder:
        copy = Path(folder) / "rehearsal.xlsx"
        shutil.copy(path, copy)
        tally = fill(copy)
        summary = recalculate(copy, Path(folder) / "out")["Resumen"]
        wrong = [(cell, want, summary[cell].value) for cell, want in expected_cells(tally).items()
                 if (abs(summary[cell].value - want) > 1e-9 if isinstance(want, float) else summary[cell].value != want)]
    if wrong:
        sys.exit(f"Resumen differs from the Python count: {wrong}")
    print(f"ok: {len(expected_cells(tally))} cells of the Resumen match the count taken in Python")


if __name__ == "__main__":
    main(sys.argv[1])
